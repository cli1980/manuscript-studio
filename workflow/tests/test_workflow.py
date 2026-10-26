"""Run: python -m unittest workflow.tests.test_workflow -v."""

import json
import contextlib
import io
import importlib.util
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from ..src.engine import DEFAULT_DEFINITION, Engine, RESULT_SCHEMA, definition_at
from ..src.runtime import CodexRuntime, Request
from ..src.storage import digest, now
from ..src.formatting import FormattingEngine
from ..src import document_adapters
from . import run_e2e
from . import monitor_performance


def result(outcome="completed", **fields):
    return dict(outcome=outcome, message="Saved operation result.", state_path=None,
                work_record=None, output_path=None, gate=None, review_path=None,
                binding_path=None) | fields


class ScriptedRuntime:
    """A deterministic test double, never presented as a model translation."""

    def __init__(self, action):
        self.action = action

    def run(self, request, on_event, on_thread):
        on_thread("test-thread")
        on_event("test.turn", {"model": request.model})
        return self.action(request)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name).resolve()
        self.source = self.base / "original.md"
        self.source.write_bytes("# 故事\n\n他回来了。\n".encode())
        self.engine = Engine.create(self.base / "run with spaces", self.source, "zh")
        self.addCleanup(lambda: self.engine.close())
        self.workspace = self.engine.workspace

    def write(self, relative, content):
        path = self.workspace / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        if not isinstance(content, str):
            content = json.dumps(content, ensure_ascii=False, indent=2) + "\n"
        path.write_bytes(content.encode())
        return relative

    def preparation(self, request):
        source = self.engine.root / "input" / self.source.name
        content = source.read_bytes()
        parts = getattr(self, "parts", [content.decode()])
        chapters = [{
            "chapter_id": "story" if len(parts) == 1 else f"chapter-{index}", "order": index, "source_number": None,
            "title": "故事", "role": "standalone", "file": "chapters/0001.md",
            "sha256": digest(part.encode()), "source_spans": [{"source_id": "original", "locator": f"Source unit {index}"}],
            "status": "verified",
        } for index, part in enumerate(parts, 1)]
        for index, chapter in enumerate(chapters, 1):
            chapter["file"] = f"chapters/{index:04d}.md"
        manifest = {
            "schema_version": 1, "work_id": "story", "revision": "r0001", "status": "verified",
            "created_at": now(), "sources": [{
                "source_id": "original", "locator": str(source), "format": "markdown",
                "identity": {"kind": "sha256", "value": digest(content)},
            }],
            "scope": {"kind": "whole_work", "chapter_ids": [c["chapter_id"] for c in chapters]},
            "chapters": chapters, "assets": [],
            "extraction": {"tool": "test", "method": "Exact fixture copy", "normalization_notes": []},
            "verification": {"performed_by": "test", "verified_at": now(),
                             "coverage_complete": True, "issues": []},
        }
        work = {
            "schema_version": 1, "work_id": "story", "record_revision": 1, "updated_at": now(),
            "metadata": {"title": "Story", "authors": [], "source_language": "zh",
                         "work_type": "fiction", "genre": None, "edition": None},
            "discovery": {"location_kind": "file", "location": str(source), "recursive": False,
                          "selection_status": "confirmed", "notes": []},
            "preparation": {"status": "ready", "current_revision": "r0001", "blockers": []},
            "workflows": {name: {"status": "not_started", "basis_revision": None, "scope": None,
                                 "operation": None, "state_locator": None, "updated_at": None}
                          for name in ("literary-editing", "literary-translation",
                                       "literary-agent-querying", "format-manuscript")},
        }
        for chapter, part in zip(chapters, parts):
            self.write("prepared/story/revisions/r0001/" + chapter["file"], part)
        self.write("prepared/story/revisions/r0001/manifest.json", manifest)
        return result(work_record=self.write("prepared/story/work.json", work))

    def prepare(self):
        self.engine.advance(ScriptedRuntime(self.preparation))
        self.engine.reply(self.engine.state()["gate"]["id"], "confirm metadata")

    def gate_receipts(self):
        return [r for r in self.engine.state()["receipts"] if r["gate"]["kind"] != "metadata_review"]

    def prepare_v2(self, request):
        response = self.preparation(request)
        path = self.workspace / response["work_record"]
        work = json.loads(path.read_text(encoding="utf-8"))
        template = self.engine.root / "skills/prepare-manuscript/references/work.example.json"
        work["schema_version"] = 2
        work["metadata"] = json.loads(template.read_text(encoding="utf-8"))["metadata"]
        work["metadata"].update(title="Story", manuscript_language="zh", genres=["Historical fiction"])
        work["metadata_provenance"] = {}
        self.write(response["work_record"], work)
        return response

    def test_metadata_review_gate_precedes_translation(self):
        self.engine.advance(ScriptedRuntime(self.preparation))
        state = self.engine.state()
        gate = state["gate"]
        self.assertEqual((state["stage"], state["status"], gate["kind"]), ("metadata", "WAITING_FOR_HUMAN", "metadata_review"))
        self.assertEqual(gate["record_revision"], 1)
        self.assertIn('/title: "Story"', gate["message"])
        self.assertIsNone(state["metadata_approval"])
        with self.assertRaisesRegex(ValueError, "WAITING_FOR_HUMAN"):
            self.engine.advance(ScriptedRuntime(self.translation))
        with self.assertRaisesRegex(ValueError, "confirm metadata"):
            self.engine.reply(gate["id"], "looks fine")
        self.assertEqual(self.engine.state()["receipts"], [])
        state = self.engine.reply(gate["id"], "confirm metadata")
        self.assertEqual((state["stage"], state["status"]), ("translate", "READY"))
        self.assertIsNone(state["next_input"])
        approval = state["metadata_approval"]
        self.assertEqual((approval["record_revision"], approval["sha256"], approval["decision_source"]),
                         (1, gate["binding"]["sha256"], "human"))
        self.assertEqual(approval["receipt_id"], state["receipts"][0]["id"])
        self.assertIn('"metadata_approval"', self.engine.prompt(state))
        self.assertEqual(self.engine.store.events()[-1]["kind"], "metadata.confirmed")
        with self.assertRaisesRegex(ValueError, "Stale"):
            self.engine.reply(gate["id"], "confirm metadata")

    def test_metadata_edits_are_versioned_and_rebind_the_gate(self):
        self.engine.advance(ScriptedRuntime(self.prepare_v2))
        first = self.engine.state()["gate"]
        self.assertIn('/genres: ["Historical fiction"] [no provenance]', first["message"])
        self.assertIn("/subtitle", first["message"].split("Unknown: ")[1])
        with self.assertRaisesRegex(ValueError, "locked"):
            self.engine.edit_metadata(first["id"], "/manuscript_language", "fr")
        with self.assertRaisesRegex(ValueError, "Stale"):
            self.engine.edit_metadata("old-gate", "/genres", ["Literary fiction"])
        state = self.engine.edit_metadata(first["id"], "/genres", ["Literary fiction"], reason="Author correction")
        gate = state["gate"]
        self.assertNotEqual(gate["id"], first["id"])
        self.assertEqual(gate["record_revision"], 2)
        self.assertIn('/genres: ["Literary fiction"] [user_edit, confirmed]', gate["message"])
        work_path = self.engine.file(state["prepared"]["work_record"])
        work = json.loads(work_path.read_text(encoding="utf-8"))
        self.assertEqual(work["metadata_provenance"]["/genres"]["rationale"], "Author correction")
        history = list((work_path.parent / "metadata-history").glob("1-*.json"))
        self.assertEqual(len(history), 1)
        self.assertEqual(json.loads(history[0].read_text(encoding="utf-8"))["metadata"]["genres"], ["Historical fiction"])
        self.assertEqual(self.engine.store.content(first["binding"]), history[0].read_bytes())
        with self.assertRaisesRegex(ValueError, "Stale"):
            self.engine.reply(first["id"], "confirm metadata")
        value = self.base / "title.json"
        value.write_text('"Corrected title"', encoding="utf-8")
        from ..src import cli
        with contextlib.redirect_stdout(io.StringIO()) as output:
            code = cli.main(["metadata", str(self.engine.root), "--gate", gate["id"], "--field", "/title",
                             "--value-file", str(value)])
        self.assertEqual(code, 0)
        self.assertIn('/title: "Corrected title"', output.getvalue())
        state = self.engine.state()
        self.assertEqual(state["gate"]["record_revision"], 3)
        self.assertEqual([e["kind"] for e in self.engine.store.events()].count("metadata.edited"), 2)
        state = self.engine.reply(state["gate"]["id"], "confirm metadata")
        self.assertEqual(state["metadata_approval"]["record_revision"], 3)
        self.assertIn("Corrected title", self.engine.prompt(state))

    def test_external_metadata_change_requires_a_refreshed_gate(self):
        self.engine.advance(ScriptedRuntime(self.prepare_v2))
        gate = self.engine.state()["gate"]
        relative = self.engine.state()["prepared"]["work_record"]
        work = json.loads(self.engine.file(relative).read_text(encoding="utf-8"))
        work["metadata"]["title"] = "Edited outside the gate"
        self.write(relative, work)
        with self.assertRaisesRegex(ValueError, "Review changed"):
            self.engine.reply(gate["id"], "confirm metadata")
        with self.assertRaisesRegex(ValueError, "Metadata changed"):
            self.engine.edit_metadata(gate["id"], "/genres", [])
        self.engine.recover(refresh_gate=True)
        state = self.engine.advance(None)
        self.assertEqual((state["stage"], state["status"]), ("metadata", "WAITING_FOR_HUMAN"))
        self.assertNotEqual(state["gate"]["id"], gate["id"])
        self.assertIn("Edited outside the gate", state["gate"]["message"])
        self.engine.reply(state["gate"]["id"], "confirm metadata")
        self.assertEqual(self.engine.state()["stage"], "translate")

    def test_editable_v2_metadata_does_not_change_prepared_source_identity(self):
        self.engine.advance(ScriptedRuntime(self.prepare_v2))
        state = self.engine.reply(self.engine.state()["gate"]["id"], "confirm metadata")
        relative = state["prepared"]["work_record"]
        work = json.loads(self.engine.file(relative).read_text(encoding="utf-8"))
        work["metadata"]["title"] = "Corrected title"
        work["metadata"]["genres"] = ["Literary fiction"]
        work["record_revision"] += 1
        self.write(relative, work)
        self.engine.verify(state)
        self.assertIn("Corrected title", self.engine.prompt(state))
        work["metadata"]["manuscript_language"] = "fr"
        self.write(relative, work)
        with self.assertRaisesRegex(ValueError, "language metadata differs"):
            self.engine.verify(state)

    def sheet(self, revision="1", resolved=False):
        return {
            "schema_version": "1.0", "sheet_id": "map-story", "revision": revision,
            "title": "Passage map", "bindings": [{"label": "Chapter", "value": "story"}],
            "decisions": [
                {"id": "approve", "command": "approve", "label": "Approve", "requires_note": False, "effect": "resolve"},
                {"id": "revise", "command": "revise", "label": "Revise", "requires_note": True, "effect": "reopen"},
            ],
            "items": [{
                "id": "PM001", "summary": "One passage: complete story",
                "details": [{"label": "Original", "value": "他回来了。"}],
                "status": "resolved" if resolved else "pending", "status_label": "Reviewed" if resolved else "Pending",
                "allowed_decisions": ["approve", "revise"], "dependencies": [],
                "response": {"decision": "approve", "note": "", "date": now()} if resolved else None,
            }],
            "gate": {"status": "awaiting-approval", "closure_required": True,
                     "final_closure": "pending", "ready_state": "not-ready"},
            "presentation": {"mode": "all", "note_label": "Feedback", "note_placeholder": "Required for revision",
                             "submit_label": "Submit"},
            "submission": {"instruction": "Review this map.", "rules": ["Do not translate before closure."],
                           "closure_command": "confirm final closure", "closure_label": "Confirm"},
        }

    def map_gate(self, revision="1", resolved=False):
        return result(
            "needs_human_input", message="Review the complete story passage boundary.",
            gate="passage_map", state_path=self.write("translation/state.md", "Passage map pending.\n"),
            binding_path=self.write("translation/map-proposal.md", "PM001: entire source story.\n"),
            review_path=self.write("translation/map-review.json", self.sheet(revision, resolved)),
        )

    def approve_map(self):
        self.engine.advance(ScriptedRuntime(lambda request: self.map_gate()))
        gate = self.engine.state()["gate"]["id"]
        self.engine.reply(gate, "PM001 approve")
        self.engine.advance(ScriptedRuntime(lambda request: self.map_gate("2", True)))
        self.engine.reply(self.engine.state()["gate"]["id"], "confirm final closure")

    def sanity_gate(self, request):
        return result(
            "needs_human_input", gate="chapter_sanity", message="He came back.\nNo textual changes proposed.",
            state_path=self.write("translation/sanity.md", "No textual changes proposed\n"),
            binding_path=self.write("translation/assembled.md", "He came back.\n"),
        )

    def translation(self, request):
        self.write("translation/glossary.md", "# Glossary\n")
        self.write("translation/voice.md", "# Voice\n")
        return result(state_path=self.write("translation/state.md", "Chapter finalized.\n"),
                      output_path=self.write("translation/translated_english/story.md", "He came back.\n"))

    def complete_translation(self):
        self.prepare()
        self.approve_map()
        self.engine.advance(ScriptedRuntime(self.sanity_gate))
        self.engine.reply(self.engine.state()["gate"]["id"], "approve chapter")
        self.engine.advance(ScriptedRuntime(self.translation))

    def test_restart_gates_acceptance_and_export_preserve_original(self):
        original = self.source.read_bytes()
        self.prepare()
        self.engine.advance(ScriptedRuntime(lambda request: self.map_gate()))
        gate_id = self.engine.state()["gate"]["id"]
        self.engine.close()
        self.engine = Engine(self.base / "run with spaces")
        self.assertEqual(self.engine.state()["gate"]["id"], gate_id)
        with self.assertRaisesRegex(ValueError, "WAITING_FOR_HUMAN"):
            self.engine.advance(ScriptedRuntime(lambda r: result()))
        with self.assertRaisesRegex(ValueError, "unresolved"):
            self.engine.reply(gate_id, "confirm final closure")
        self.engine.reply(gate_id, "PM001 approve")
        with self.assertRaisesRegex(ValueError, "Stale"):
            self.engine.reply(gate_id, "PM001 approve")
        self.engine.advance(ScriptedRuntime(lambda request: self.map_gate("2", True)))
        self.engine.reply(self.engine.state()["gate"]["id"], "confirm final closure")
        self.engine.advance(ScriptedRuntime(self.sanity_gate))
        self.engine.reply(self.engine.state()["gate"]["id"], "approve chapter")
        self.engine.advance(ScriptedRuntime(self.translation))
        self.assertEqual(self.engine.state()["status"], "WAITING_FOR_ACCEPTANCE")
        with self.assertRaisesRegex(ValueError, "Stale"):
            self.engine.accept("incorrect-acceptance")
        self.engine.accept(self.engine.state()["acceptance_id"])
        state = self.engine.advance(None)
        self.assertEqual(state["status"], "COMPLETED")
        self.assertEqual((Path(state["export_path"]) / "0001.md").read_text(), "He came back.\n")
        self.assertEqual(self.source.read_bytes(), original)
        self.assertTrue(any(event["kind"] == "review.received" for event in self.engine.store.events()))

    def test_cannot_complete_translation_without_domain_gates(self):
        self.prepare()
        with self.assertRaisesRegex(ValueError, "map closure"):
            self.engine.advance(ScriptedRuntime(self.translation))
        self.assertEqual(self.engine.state()["status"], "FAILED")
        self.assertEqual(self.engine.state()["outputs"], [])
        self.engine.recover()
        self.assertEqual(self.engine.state()["stage"], "translate")
        self.assertIsNone(self.engine.state()["thread_id"])

    def test_second_chapter_has_new_context_and_cannot_inherit_first_chapter_gates(self):
        self.engine.close()
        self.parts = ["# 第一章\n\n他回来了。\n", "# 第二章\n\n她走了。\n"]
        self.source.write_bytes("".join(self.parts).encode())
        self.engine = Engine.create(self.base / "two-chapters", self.source, "zh")
        self.workspace = self.engine.workspace
        self.complete_translation()
        state = self.engine.state()
        self.assertEqual(state["status"], "READY")
        self.assertEqual(state["chapter_index"], 1)
        self.assertIsNone(state["thread_id"])
        self.assertIsNone(state["map_approval"])
        self.assertIsNone(state["sanity_approval"])
        self.assertEqual(len(state["outputs"]), 1)
        self.assertEqual(set(state["continuity"]), {"glossary", "voice"})
        with self.assertRaisesRegex(ValueError, "map closure"):
            self.engine.advance(ScriptedRuntime(self.translation))

    def test_malformed_result_is_persisted_as_failure(self):
        with self.assertRaisesRegex(ValueError, "Invalid runtime result"):
            self.engine.advance(ScriptedRuntime(lambda request: {"outcome": "completed"}))
        self.assertEqual(self.engine.state()["status"], "FAILED")
        self.assertEqual(self.engine.state()["stage"], "prepare")

    def test_empty_sanity_diff_still_requires_explicit_approval(self):
        self.prepare()
        self.approve_map()
        self.engine.advance(ScriptedRuntime(self.sanity_gate))
        self.engine.reply(self.engine.state()["gate"]["id"], "continue")
        self.assertIsNone(self.engine.state()["sanity_approval"])
        with self.assertRaisesRegex(ValueError, "chapter approval"):
            self.engine.advance(ScriptedRuntime(self.translation))

    def test_agent_cannot_claim_resolved_items_without_human_receipts(self):
        self.prepare()
        with self.assertRaisesRegex(ValueError, "no matching human"):
            self.engine.advance(ScriptedRuntime(lambda request: self.map_gate("1", True)))
        self.assertIsNone(self.engine.state()["map_approval"])

    def test_changed_review_and_changed_source_block_responses(self):
        self.prepare()
        self.engine.advance(ScriptedRuntime(lambda request: self.map_gate()))
        gate_id = self.engine.state()["gate"]["id"]
        self.write("translation/map-proposal.md", "Changed boundary\n")
        with self.assertRaisesRegex(ValueError, "Review changed"):
            self.engine.reply(gate_id, "PM001 approve")
        self.write("prepared/story/revisions/r0001/chapters/0001.md", "Tampered original\n")
        with self.assertRaisesRegex(ValueError, "Prepared source changed"):
            self.engine.recover(refresh_gate=True)

    def test_wrong_revision_envelope_is_rejected(self):
        self.prepare()
        self.engine.advance(ScriptedRuntime(lambda request: self.map_gate()))
        sheet = self.sheet()
        payload = "\n".join([sheet["submission"]["instruction"], "Approval sheet: map-story",
                             "Revision: old", *sheet["submission"]["rules"], "", "PM001 approve"])
        with self.assertRaises(ValueError):
            self.engine.reply(self.engine.state()["gate"]["id"], payload)
        self.assertEqual(self.gate_receipts(), [])

    def test_changed_proposal_cannot_reuse_old_disposition(self):
        self.prepare()
        self.engine.advance(ScriptedRuntime(lambda request: self.map_gate()))
        self.engine.reply(self.engine.state()["gate"]["id"], "PM001 approve")

        def changed(request):
            output = self.map_gate("2", True)
            sheet = self.sheet("2", True)
            sheet["items"][0]["details"][0]["value"] = "A different boundary"
            self.write(output["review_path"], sheet)
            return output

        with self.assertRaisesRegex(ValueError, "no matching human"):
            self.engine.advance(ScriptedRuntime(changed))

    def test_source_failure_and_path_escape_do_not_promote(self):
        def invalid(request):
            output = self.preparation(request)
            path = self.workspace / output["work_record"]
            work = json.loads(path.read_text())
            work["preparation"]["blockers"] = ["Missing page"]
            path.write_text(json.dumps(work), encoding="utf-8")
            return output

        with self.assertRaisesRegex(ValueError, "blockers"):
            self.engine.advance(ScriptedRuntime(invalid))
        self.assertEqual(self.engine.state()["stage"], "prepare")
        self.engine.recover()
        with self.assertRaisesRegex(ValueError, "escapes"):
            self.engine.advance(ScriptedRuntime(lambda r: result(work_record="../outside.json")))

    def test_failure_keeps_receipt_and_restart_requires_explicit_recovery(self):
        self.prepare()
        self.engine.advance(ScriptedRuntime(lambda request: self.map_gate()))
        self.engine.reply(self.engine.state()["gate"]["id"], "PM001 approve")
        receipt = self.engine.state()["next_input"]["id"]

        def fail(request):
            raise RuntimeError("Interrupted provider")

        with self.assertRaisesRegex(RuntimeError, "Interrupted"):
            self.engine.advance(ScriptedRuntime(fail))
        self.engine.recover()
        self.assertEqual(self.engine.state()["next_input"]["id"], receipt)
        self.assertEqual(len(self.engine.state()["receipts"]), 2)

    def test_export_conflict_is_not_overwritten(self):
        self.complete_translation()
        self.engine.accept(self.engine.state()["acceptance_id"])
        export = self.engine.root / "export"
        export.mkdir()
        (export / "0001.md").write_bytes(b"User file")
        with self.assertRaisesRegex(ValueError, "refusing to overwrite"):
            self.engine.advance(None)
        self.assertEqual((export / "0001.md").read_bytes(), b"User file")

    def test_run_lock_excludes_another_process_handle(self):
        second = Engine(self.engine.root)
        try:
            with self.engine.store.locked():
                with self.assertRaisesRegex(ValueError, "Another process"):
                    with second.store.locked():
                        self.fail("Acquired an already-held lock")
        finally:
            second.close()

    def test_reject_unsupported_language_or_incomplete_skill_bundle(self):
        config = json.loads(DEFAULT_DEFINITION.read_text())
        path = self.base / "definition.json"
        config["target_language"] = "fr-FR"
        path.write_text(json.dumps(config))
        with self.assertRaisesRegex(ValueError, "US English"):
            definition_at(path)
        config["target_language"] = "en-US"
        config["agents"]["translate"]["skills"].remove("review-approval-sheet")
        path.write_text(json.dumps(config))
        with self.assertRaisesRegex(ValueError, "bundle"):
            definition_at(path)

    def e2e_case(self):
        spec = {
            "schema_version": 1, "case_id": "fixture-e2e", "revision": 1,
            "task": "literary-translation-workflow",
            "input": {
                "manuscript": {"file": str(self.source), "format": "markdown", "sha256": digest(self.source.read_bytes())},
                "source_language": "zh", "target_language": "en-US", "scope": "whole_work",
                "workflow_definition": str(DEFAULT_DEFINITION), "task_specification": "task.md",
                "runtime": {"adapter": "codex-app-server", "model": None, "human_responses": "interactive"},
            },
            "source_inventory": {"title": "Synthetic fixture"},
            "expected_output": {
                "stage_sequence": ["prepare", "metadata", "translate", "accept", "export"],
                "required_files": ["state.sqlite3", "export/manifest.json", "export/glossary.md", "export/voice.md"],
            },
            "checks": [
                {"id": name, "evaluation": method, "expectation": f"Verify {name} on the synthetic fixture."}
                for name, method in (
                    ("source-integrity", "deterministic"), ("coverage", "source-comparison"),
                    ("domain-gates", "receipts-and-domain-review"), ("recovery", "interactive-scenario"),
                    ("stale-response", "interactive-scenario"), ("continuity", "domain-review"),
                    ("export-provenance", "deterministic"), ("fidelity-routing", "trace-and-domain-review"),
                )
            ],
        }
        (self.base / "task.md").write_text("Translate the complete synthetic fixture.", encoding="utf-8")
        path = self.base / "case.json"
        path.write_text(json.dumps(spec), encoding="utf-8")
        case = run_e2e.load_case(path)
        run_e2e.bind_case(self.engine, case, {"executable": "test-only", "version": "scripted fixture"})
        return case

    def test_e2e_launcher_requires_an_explicit_case(self):
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as stopped:
            run_e2e.parser().parse_args([])
        self.assertEqual(stopped.exception.code, 2)

    def test_e2e_launcher_interactive_loop_and_honest_report(self):
        case = self.e2e_case()
        actions = iter([
            self.preparation, lambda r: self.map_gate(),
            lambda r: self.map_gate("2", True), self.sanity_gate, self.translation,
        ])
        prompts = []

        def action(request):
            prompts.append(request.prompt)
            return next(actions)(request)

        replies = iter(["confirm metadata", ".", "PM001 approve", ".", "confirm final closure", ".",
                        "approve chapter", ".", "accept", "skip", "skip", "skip", "skip"])
        runtime = run_e2e.CaseRuntime(ScriptedRuntime(action), case["task"])
        with contextlib.redirect_stdout(io.StringIO()):
            code = run_e2e.drive(self.engine, case, runtime, input_fn=lambda _: next(replies))
        self.assertEqual(code, 2)  # Execution completed, manual checks not claimed as passed.
        self.assertTrue(all("Translate the complete synthetic fixture." in prompt for prompt in prompts))
        report = json.loads((self.engine.root / "e2e-report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["workflow_status"], "COMPLETED")
        self.assertEqual(report["e2e_status"], "needs-review")
        checks = {c["id"]: c for c in report["checks"]}
        for name in ("recovery", "stale-response", "source-integrity", "export-provenance"):
            self.assertEqual(checks[name]["status"], "passed")
        self.assertEqual(checks["domain-gates"]["status"], "not-checked")
        self.assertEqual(report["literary_quality"]["status"], "not-checked")
        self.assertEqual((self.engine.root / "candidate.en.txt").read_text(), "He came back.\n")
        self.assertEqual(len(self.engine.state()["receipts"]), 4)
        self.assertEqual(report["automated_checks"]["stage-order"]["evidence"],
                         ["prepare", "metadata", "translate", "accept", "export"])
        # A completed run can finish its human evidence review without another model call.
        evidence = iter(["pass", "Synthetic test evidence"] * 4)
        with contextlib.redirect_stdout(io.StringIO()):
            code = run_e2e.drive(self.engine, case, None, input_fn=lambda _: next(evidence))
        self.assertEqual(code, 0)
        report = json.loads((self.engine.root / "e2e-report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["e2e_status"], "passed")
        self.assertEqual(report["literary_quality"]["status"], "not-checked")

    def test_e2e_launcher_eof_preserves_gate_and_can_bind_resume(self):
        case = self.e2e_case()
        self.prepare()
        self.engine.advance(ScriptedRuntime(lambda r: self.map_gate()))
        gate_id = self.engine.state()["gate"]["id"]

        def eof(_):
            raise EOFError

        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(run_e2e.drive(self.engine, case, None, input_fn=eof), 2)
        self.assertEqual(self.engine.state()["gate"]["id"], gate_id)
        self.assertEqual(self.gate_receipts(), [])
        run_e2e.bind_case(self.engine, case, {"version": "test"}, resume=True)
        changed = dict(case, task_sha256="changed")
        with self.assertRaisesRegex(ValueError, "Changed task_sha256"):
            run_e2e.bind_case(self.engine, changed, {}, resume=True)

    def test_e2e_report_detects_corrupt_export_and_does_not_write_candidate(self):
        case = self.e2e_case()
        self.complete_translation()
        self.engine.accept(self.engine.state()["acceptance_id"])
        self.engine.advance(None)
        (self.engine.root / "export/0001.md").write_bytes(b"Unaccepted replacement")
        report = run_e2e.write_report(self.engine, case)
        self.assertEqual(report["e2e_status"], "failed")
        self.assertEqual(report["automated_checks"]["export-provenance"]["status"], "failed")
        run_e2e.collect_candidate(self.engine, report)
        self.assertFalse((self.engine.root / "candidate.en.txt").exists())

    def test_e2e_case_check_has_no_run_or_runtime_side_effects(self):
        case = self.e2e_case()
        unused = self.base / "unused-run"
        with contextlib.redirect_stdout(io.StringIO()):
            code = run_e2e.main(["--case", case["case_path"], "--check-case", "--run-dir", str(unused)])
        self.assertEqual(code, 0)
        self.assertFalse(unused.exists())
        self.source.write_bytes(b"Changed source")
        with self.assertRaisesRegex(ValueError, "manuscript hash"):
            run_e2e.load_case(Path(case["case_path"]))

    def test_recommended_approvals_complete_without_claiming_human_review(self):
        case = self.e2e_case()
        actions = iter([
            (self.preparation, "accept", "confirm metadata"),
            (lambda r: self.map_gate(), "accept", "PM001 approve"),
            (lambda r: self.map_gate("2", True), "accept", "confirm final closure"),
            (self.sanity_gate, "accept", "approve chapter"),
            (self.translation, "accept", "accept"),
        ])

        def action(request):
            operation, recommended, response = next(actions)
            self.assertIn("approval_recommendation", request.output_schema["required"])
            return operation(request) | {"approval_recommendation": {
                "action": recommended, "response": response, "reason": "Synthetic fixture recommendation",
            }}

        runtime = run_e2e.CaseRuntime(ScriptedRuntime(action), case["task"], auto_approve_recommended=True)
        with contextlib.redirect_stdout(io.StringIO()):
            code = run_e2e.drive(self.engine, case, runtime, auto_approve_recommended=True,
                                 input_fn=lambda _: self.fail("Automation must not read stdin."))
        self.assertEqual(code, 2)
        state = self.engine.state()
        self.assertEqual(state["status"], "COMPLETED")
        self.assertTrue(all(r["decision_source"] == "delegated-automation" for r in state["receipts"]))
        self.assertEqual(state["acceptance"]["decision_source"], "delegated-automation")
        report = json.loads((self.engine.root / "e2e-report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["e2e_status"], "needs-review")
        self.assertEqual(report["delegated_decisions"], 4)
        self.assertEqual(state["metadata_approval"]["decision_source"], "delegated-automation")
        self.assertEqual(report["literary_quality"]["status"], "not-checked")

    def test_automation_does_not_confirm_unrecommended_or_edited_metadata(self):
        case = self.e2e_case()
        for recommendation, edit, reason in [
            ({"action": "needs_review", "response": "", "reason": "Genre needs the author's judgment"}, False,
             "Recommendation requires review"),
            ({"action": "accept", "response": "confirm metadata", "reason": "Synthetic fixture"}, True,
             "Metadata changed after the recommendation"),
        ]:
            with self.subTest(edit=edit):
                self.engine.close()
                self.engine = Engine.create(self.base / f"automation-{edit}", self.source, "zh")
                self.workspace = self.engine.workspace
                run_e2e.bind_case(self.engine, case, {"executable": "test-only", "version": "scripted fixture"})
                runtime = run_e2e.CaseRuntime(
                    ScriptedRuntime(lambda r: self.prepare_v2(r) | {"approval_recommendation": recommendation}),
                    case["task"], auto_approve_recommended=True,
                )
                self.engine.advance(runtime)
                if edit:
                    self.engine.edit_metadata(self.engine.state()["gate"]["id"], "/genres", [])
                with contextlib.redirect_stdout(io.StringIO()) as output:
                    code = run_e2e.drive(self.engine, case, runtime, auto_approve_recommended=True,
                                         input_fn=lambda _: self.fail("Must stop without reading stdin."))
                self.assertEqual(code, 2)
                self.assertIn(reason, output.getvalue())
                state = self.engine.state()
                self.assertEqual((state["stage"], state["status"]), ("metadata", "WAITING_FOR_HUMAN"))
                self.assertEqual(state["receipts"], [])

    def test_automation_stops_on_review_or_nonaccepting_recommendation(self):
        case = self.e2e_case()
        self.prepare()
        for recommendation in [
            {"action": "needs_review", "response": "", "reason": "Ambiguous source needs user input"},
            {"action": "accept", "response": "PM001 revise: move boundary", "reason": "Invalid accepting action"},
        ]:
            with self.subTest(recommendation=recommendation):
                if self.engine.state()["status"] == "WAITING_FOR_HUMAN":
                    self.engine.recover(refresh_gate=True)
                runtime = run_e2e.CaseRuntime(
                    ScriptedRuntime(lambda r: self.map_gate() | {"approval_recommendation": recommendation}),
                    case["task"], auto_approve_recommended=True,
                )
                with contextlib.redirect_stdout(io.StringIO()):
                    code = run_e2e.drive(self.engine, case, runtime, auto_approve_recommended=True,
                                         input_fn=lambda _: self.fail("Must stop without reading stdin."))
                self.assertEqual(code, 2)
                self.assertEqual(self.engine.state()["status"], "WAITING_FOR_HUMAN")
                self.assertEqual(self.gate_receipts(), [])

    def formatting_job(self):
        self.complete_translation()
        self.engine.accept(self.engine.state()["acceptance_id"])
        self.engine.advance(None)
        adapter = self.base / "documents"
        adapter.mkdir()
        (adapter / "SKILL.md").write_text("# Synthetic Documents adapter for tests\n", encoding="utf-8")
        job = FormattingEngine.create_from_run(self.engine, documents_skill=adapter)
        self.addCleanup(job.close)
        return job, adapter

    def format_gate(self, job, *, revision="1", resolved=False):
        folder = job.workspace / "manuscript-formatting"
        folder.mkdir(exist_ok=True)
        (folder / "format-spec.md").write_text("# Proposed format\n- Status: proposed\nReadable fiction edition.\n", encoding="utf-8")
        (folder / "proposal.md").write_text("# Proposed format\nReadable fiction edition.\n", encoding="utf-8")
        sheet = self.sheet(revision, resolved)
        sheet["sheet_id"] = "format-fixture"
        sheet["items"][0]["id"] = "FM001"
        sheet["items"][0]["summary"] = "Use the proposed reader-facing edition."
        sheet["items"][0]["details"] = [{"label": "Design", "value": "Readable fiction edition."}]
        (folder / "format-review.json").write_text(json.dumps(sheet), encoding="utf-8")
        output = result("needs_human_input", gate="format_spec",
                      state_path="manuscript-formatting/format-spec.md",
                      binding_path="manuscript-formatting/proposal.md",
                      review_path="manuscript-formatting/format-review.json", qa_path=None)
        if job.state().get("backend"):
            (folder / "format-settings.json").write_text(json.dumps(document_adapters.DEFAULT_SETTINGS), encoding="utf-8")
            output["settings_path"] = "manuscript-formatting/format-settings.json"
        return output

    def approve_format(self, job):
        job.advance(ScriptedRuntime(lambda r: self.format_gate(job)))
        job.reply(job.state()["gate"]["id"], "FM001 approve")
        job.advance(ScriptedRuntime(lambda r: self.format_gate(job, revision="2", resolved=True)))
        job.reply(job.state()["gate"]["id"], "confirm final closure")

    def format_result(self, job, *, text="He came back.", pages=True):
        from zipfile import ZipFile
        import struct
        import zlib
        from xml.sax.saxutils import escape
        folder = job.workspace / "manuscript-formatting/outputs/docx"
        folder.mkdir(parents=True, exist_ok=True)
        document = folder / "manuscript.docx"
        with ZipFile(document, "w") as archive:
            archive.writestr("[Content_Types].xml",
                             '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                             '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                             '</Types>')
            archive.writestr("_rels/.rels",
                             '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                             '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                             '</Relationships>')
            archive.writestr("word/document.xml", (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                f'<w:body><w:p><w:r><w:t>{escape(text)}</w:t></w:r></w:p></w:body></w:document>'
            ))
            archive.writestr("word/styles.xml",
                             '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        render = folder / "render"
        render.mkdir(exist_ok=True)

        def chunk(kind, data):
            return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

        png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
               + chunk(b"IDAT", zlib.compress(b"\0\xff\xff\xff")) + chunk(b"IEND", b""))
        (render / "page-1.png").write_bytes(png)
        qa = {
            "schema_version": 1, "input_manifest_sha256": job.state()["input_manifest_sha256"],
            "approved_spec_sha256": (job.state().get("format_approval") or {}).get("sha256", "0" * 64),
            "docx_sha256": digest(document.read_bytes()), "adapter_skill_sha256": job.state()["adapter_sha256"],
            "visual_review_passed": True, "specification_matches": True, "page_count": 1,
            "content_paragraph_indexes": [0],
            "pages": [{"page": 1, "image": "manuscript-formatting/outputs/docx/render/page-1.png",
                       "sha256": digest(png), "reviewed": True, "findings": []}] if pages else [],
        }
        (folder / "qa.json").write_text(json.dumps(qa), encoding="utf-8")
        # Fixtures exercise the contract, not a real render or human visual assessment.
        return result(state_path="manuscript-formatting/format-spec.md",
                      output_path="manuscript-formatting/outputs/docx/manuscript.docx",
                      qa_path="manuscript-formatting/outputs/docx/qa.json")

    def test_formatter_requires_closure_and_preserves_accepted_inputs(self):
        job, adapter = self.formatting_job()
        accepted = self.engine.state()["acceptance"]
        original = self.source.read_bytes()
        self.approve_format(job)
        job.advance(ScriptedRuntime(lambda r: self.format_result(job)))
        state = job.state()
        self.assertEqual(state["status"], "COMPLETED")
        self.assertTrue(Path(state["document_path"]).is_file())
        self.assertTrue((job.root / "export/manifest.json").is_file())
        self.assertEqual(self.engine.state()["status"], "COMPLETED")
        self.assertEqual(self.engine.state()["acceptance"], accepted)
        self.assertEqual(self.source.read_bytes(), original)
        reopened = Engine.open(job.root)
        self.addCleanup(reopened.close)
        self.assertIsInstance(reopened, FormattingEngine)
        reused = FormattingEngine.create_from_run(self.engine, documents_skill=adapter)
        self.addCleanup(reused.close)
        self.assertEqual(reused.root, job.root)
        self.assertEqual(len(self.engine.state()["format_jobs"]), 1)
        separate = FormattingEngine.create_from_run(self.engine, documents_skill=adapter, new_edition=True)
        self.addCleanup(separate.close)
        self.assertNotEqual(separate.root, job.root)

    def test_formatter_does_not_publish_before_approval(self):
        job, _ = self.formatting_job()
        self.format_gate(job)
        with self.assertRaisesRegex(ValueError, "not been approved"):
            job.advance(ScriptedRuntime(lambda r: self.format_result(job)))
        self.assertFalse((job.root / "export").exists())

    def test_formatter_rejects_prose_changes_even_with_successful_qa_claim(self):
        job, _ = self.formatting_job()
        self.approve_format(job)
        with self.assertRaisesRegex(ValueError, "changes accepted manuscript text"):
            job.advance(ScriptedRuntime(lambda r: self.format_result(job, text="He returned triumphantly.")))
        self.assertEqual(job.state()["outputs"], [])
        self.assertFalse((job.root / "export").exists())

    def test_formatter_rejects_missing_page_evidence_and_stale_specification(self):
        job, _ = self.formatting_job()
        job.advance(ScriptedRuntime(lambda r: self.format_gate(job)))
        (job.workspace / "manuscript-formatting/proposal.md").write_text("Changed proposal", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Review changed"):
            job.reply(job.state()["gate"]["id"], "FM001 approve")
        job.recover(refresh_gate=True)
        self.approve_format(job)
        with self.assertRaisesRegex(ValueError, "page coverage"):
            job.advance(ScriptedRuntime(lambda r: self.format_result(job, pages=False)))

    def test_formatter_rejects_incomplete_translation_and_missing_adapter(self):
        adapter = self.base / "documents"
        adapter.mkdir()
        (adapter / "SKILL.md").write_text("# Test adapter", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "completed translation"):
            FormattingEngine.create_from_run(self.engine, documents_skill=adapter)
        with self.assertRaisesRegex(ValueError, "unavailable"):
            FormattingEngine.create_from_run(self.engine, documents_skill=self.base / "missing")

    def test_optional_formatter_block_does_not_invalidate_translation_result(self):
        case = self.e2e_case()
        job, adapter = self.formatting_job()
        run_e2e.write_report(self.engine, case)
        runtime = ScriptedRuntime(lambda request: result(
            "blocked", message="blocked-adapter: renderer unavailable", qa_path=None,
        ))
        with contextlib.redirect_stdout(io.StringIO()):
            code = run_e2e.format_document(
                self.engine, runtime, documents_skill=adapter, auto_approve_recommended=True,
            )
        self.assertEqual(code, 2)
        self.assertEqual(self.engine.state()["status"], "COMPLETED")
        self.assertEqual(job.state()["status"], "BLOCKED")
        report = json.loads((self.engine.root / "e2e-report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["workflow_status"], "COMPLETED")
        self.assertEqual(report["formatting_jobs"][job.state()["id"]]["status"], "BLOCKED")
        self.assertFalse((job.root / "export/manuscript.docx").exists())

    def test_native_formatter_uses_approved_settings_and_waits_for_visual_qa(self):
        self.formatting_job()  # Complete the parent translation using synthetic fixtures.
        with patch.object(document_adapters, "preflight", return_value={"renderer": "word", "version": "test"}):
            job = FormattingEngine.create_from_run(self.engine, renderer="word")
        self.addCleanup(job.close)
        self.approve_format(job)
        self.assertIsNotNone(job.state()["approved_settings"])

        def render(markdown, settings, directory, renderer):
            self.assertEqual(markdown, "He came back.\n")
            self.assertEqual(renderer, "word")
            directory.mkdir(parents=True)
            self.format_result(job)
            folder = job.workspace / "manuscript-formatting/outputs/docx"
            pdf = folder / "fixture.pdf"
            pdf.write_bytes(b"%PDF synthetic fixture; not an actual renderer result")
            image = folder / "render/page-1.png"
            return {
                "docx": str(folder / "manuscript.docx"), "pdf": str(pdf),
                "pages": [{"page": 1, "path": str(image), "sha256": digest(image.read_bytes())}],
                "renderer": "word", "renderer_version": "fixture", "builder": "python-docx",
                "builder_version": "fixture", "toc_pages": {}, "layout_passes": 1,
            }

        with patch.object(document_adapters, "build_and_render", side_effect=render) as renderer:
            job.advance(None)
            self.assertEqual(renderer.call_count, 1)
        self.assertEqual(job.state()["status"], "READY")
        generated = job.state()["generated"]
        qa_path = job.file(generated["qa_path"])
        qa = json.loads(qa_path.read_text(encoding="utf-8"))
        self.assertFalse(qa["visual_review_passed"])
        self.assertFalse((job.root / "export/manuscript.docx").exists())
        self.assertIn("- Status: approved", job.file("manuscript-formatting/format-spec.md").read_text())
        completion = result(
            state_path="manuscript-formatting/format-spec.md", output_path=generated["document"]["path"],
            qa_path=generated["qa_path"], settings_path="manuscript-formatting/format-settings.json",
        )
        with self.assertRaisesRegex(ValueError, "Visual/specification QA"):
            job.advance(ScriptedRuntime(lambda r: completion))
        job.recover()
        qa["visual_review_passed"] = qa["specification_matches"] = True
        for page in qa["pages"]:
            page.update(reviewed=True, findings=[])
        qa_path.write_text(json.dumps(qa), encoding="utf-8")
        job.advance(ScriptedRuntime(lambda r: completion))
        self.assertEqual(job.state()["status"], "COMPLETED")
        self.assertTrue(Path(job.state()["document_path"]).is_file())

    def test_native_settings_are_revision_bound_and_renderer_switch_creates_an_edition(self):
        self.formatting_job()
        with patch.object(document_adapters, "preflight", side_effect=lambda name: {"renderer": name, "version": "test"}):
            word = FormattingEngine.create_from_run(self.engine, renderer="word")
            cloud = FormattingEngine.create_from_run(self.engine, renderer="libreoffice")
        self.addCleanup(word.close)
        self.addCleanup(cloud.close)
        self.assertNotEqual(word.root, cloud.root)
        word.advance(ScriptedRuntime(lambda r: self.format_gate(word)))
        settings = word.file("manuscript-formatting/format-settings.json")
        changed = dict(document_adapters.DEFAULT_SETTINGS, font_size=12)
        settings.write_text(json.dumps(changed), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Review changed"):
            word.reply(word.state()["gate"]["id"], "FM001 approve")


class DocumentAdapterTests(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec("docx"), "Install workflow/requirements-docx.txt for builder tests.")
    def test_builder_preserves_text_emphasis_headings_and_page_fields(self):
        from docx import Document
        from ..src.formatting import docx_paragraphs
        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "manuscript.docx"
            source = "# A Story\n\n## 1\n\nHe **came** *back*.\n\n## 2\n\nShe stayed.\n"
            settings = dict(document_adapters.DEFAULT_SETTINGS, include_toc=False)
            document_adapters.PythonDocxBuilder().build(source, settings, target, {})
            paragraphs, folios = docx_paragraphs(target)
            self.assertEqual([p["text"] for p in paragraphs], ["A Story", "1", "He came back.", "2", "She stayed."])
            self.assertTrue(paragraphs[0]["title"])
            self.assertTrue(paragraphs[1]["heading"])
            self.assertTrue(folios)
            runs = Document(target).paragraphs[2].runs
            self.assertTrue(next(r for r in runs if r.text == "came").bold)
            self.assertTrue(next(r for r in runs if r.text == "back").italic)
            from docx.oxml.ns import qn
            document = Document(target)
            for name in ("Title", "Heading 1"):
                style = document.styles[name]
                self.assertIsNone(style.element.find(qn("w:pPr") + "/" + qn("w:pBdr")))
                fonts = style.element.find(qn("w:rPr") + "/" + qn("w:rFonts"))
                self.assertFalse(any("theme" in key.lower() for key in fonts.attrib))
                self.assertEqual(fonts.get(qn("w:ascii")), settings["body_font"])
            self.assertEqual(document.paragraphs[2].paragraph_format.first_line_indent, 0)
            settings["indent_opening_paragraphs"] = True
            document_adapters.PythonDocxBuilder().build(source, settings, target, {})
            self.assertIsNone(Document(target).paragraphs[2].paragraph_format.first_line_indent)

    def test_bad_settings_and_unsupported_content_are_not_silently_converted(self):
        with self.assertRaisesRegex(ValueError, "text width"):
            document_adapters.validate_settings(dict(document_adapters.DEFAULT_SETTINGS, margin_inches=2.5))
        with self.assertRaisesRegex(ValueError, "Unsupported manuscript block"):
            document_adapters.content_blocks("# Title\n\n![illustration](image.png)\n")

    def test_cloud_renderer_uses_isolated_libreoffice_profile(self):
        from types import SimpleNamespace
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "book.docx"
            source.write_bytes(b"test input")

            def run(command, **kwargs):
                if "--convert-to" in command:
                    self.assertIn("--headless", command)
                    self.assertTrue(any(arg.startswith("-env:UserInstallation=file:") for arg in command))
                    self.assertIn("pdf:writer_pdf_Export", command)
                    (root / "book.pdf").write_bytes(b"%PDF fixture")
                return SimpleNamespace(returncode=0, stdout="LibreOffice test", stderr="")

            with patch.object(document_adapters.shutil, "which", return_value="/usr/bin/soffice"), \
                    patch.object(document_adapters.subprocess, "run", side_effect=run):
                result = document_adapters.LibreOfficeRenderer().render(source, root)
            self.assertEqual(result.renderer, "libreoffice")
            self.assertEqual(result.pdf.name, "book.pdf")


class PerformanceTests(unittest.TestCase):
    def test_approval_submission_and_followup_work_are_separate(self):
        from datetime import datetime, timezone, timedelta
        base = datetime(2026, 1, 1, tzinfo=timezone.utc)
        data = [
            (0, "execution.started", {"stage": "translate"}),
            (10, "execution.result", {"outcome": "needs_human_input", "gate": "passage_map"}),
            (10.2, "e2e.delegated_approval", {"gate_id": "g1"}),
            (10.5, "review.received", {"receipt_id": "r1", "gate_id": "g1"}),
            (11, "execution.started", {"stage": "translate"}),
            (31, "execution.result", {"outcome": "needs_human_input", "gate": "passage_map"}),
            (31.2, "e2e.delegated_approval", {"gate_id": "g2"}),
            (31.4, "review.received", {"receipt_id": "r2", "gate_id": "g2"}),
            (32, "execution.started", {"stage": "translate"}),
            (92, "execution.result", {"outcome": "needs_human_input", "gate": "continuation"}),
        ]
        events = [{"created_at": (base + timedelta(seconds=second)).isoformat(),
                   "kind": kind, "payload": payload} for second, kind, payload in data]
        state = {
            "id": "test", "status": "WAITING_FOR_HUMAN", "stage": "translate",
            "outputs": [], "chapters": [{}],
            "receipts": [
                {"id": "r1", "decision_source": "delegated-automation",
                 "gate": {"kind": "passage_map"},
                 "parsed": {"type": "dispositions", "dispositions": [{}, {}]}},
                {"id": "r2", "decision_source": "delegated-automation",
                 "gate": {"kind": "passage_map"}, "parsed": {"type": "final-closure"}},
            ],
        }
        report = monitor_performance.analyze(events, state, base + timedelta(seconds=93))
        self.assertEqual(report["work_groups"]["map_decision_recording"]["wall_seconds"], 20)
        self.assertEqual(report["work_groups"]["drafting_and_source_review"]["wall_seconds"], 60)
        self.assertEqual(report["delegated_approval_timing"]["submission_seconds"], .5)
        self.assertEqual(report["delegated_approval_timing"]["result_to_receipt_seconds"], .9)
        self.assertEqual(report["delegated_approval_timing"]["batches"][0]["item_count"], 2)
        self.assertEqual(monitor_performance.work_category(
            "translate", "translation_issue", "feedback", "continuation"),
            "issue_application_and_finalization_mixed")
        self.assertEqual(monitor_performance.work_category(
            "translate", "continuation", "feedback", "chapter_sanity"), "chapter_sanity_review")

    def test_timing_unions_tools_and_deduplicates_cumulative_usage(self):
        from datetime import datetime, timezone, timedelta
        base = datetime(2026, 1, 1, tzinfo=timezone.utc)
        events = []

        def event(second, kind, payload):
            events.append({"created_at": (base + timedelta(seconds=second)).isoformat(),
                           "kind": kind, "payload": payload})

        def usage(second, turn_id, inputs, cached, outputs, last_input):
            event(second, "thread/tokenUsage/updated", {
                "threadId": "thread", "turnId": turn_id,
                "tokenUsage": {"total": {"inputTokens": inputs, "cachedInputTokens": cached,
                                        "outputTokens": outputs, "reasoningOutputTokens": 3},
                               "last": {"inputTokens": last_input}},
            })

        event(0, "execution.started", {"stage": "translate"})
        event(.1, "execution.thread", {"thread_id": "thread"})
        event(.2, "thread/started", {"thread": {"id": "thread", "model": "test-model"}})
        event(1, "turn/started", {"turn": {"id": "turn-1"}})
        event(2, "item/started", {"turnId": "turn-1", "item": {"id": "command", "type": "commandExecution"}})
        event(3, "item/started", {"turnId": "turn-1", "item": {"id": "nested", "type": "fileChange"}})
        usage(3.5, "turn-1", 100, 80, 10, 100)
        event(4, "item/completed", {"turnId": "turn-1", "item": {"id": "nested", "type": "fileChange", "changes": []}})
        usage(4.1, "turn-1", 100, 80, 10, 100)
        event(6, "item/completed", {"turnId": "turn-1", "item": {
            "id": "command", "type": "commandExecution", "exitCode": 0, "aggregatedOutput": "fixture",
        }})
        usage(8, "turn-1", 300, 230, 30, 200)
        event(9, "turn/completed", {"turn": {"id": "turn-1"}})
        event(10, "execution.result", {"outcome": "needs_human_input", "gate": "continuation"})
        event(12, "execution.started", {"stage": "translate"})
        event(12.1, "execution.thread", {"thread_id": "thread"})
        usage(12.2, "turn-1", 300, 230, 30, 200)  # Resume notification repeats earlier usage.
        event(13, "turn/started", {"turn": {"id": "turn-2"}})
        usage(14, "turn-2", 450, 330, 45, 150)
        state = {"id": "test", "status": "RUNNING", "stage": "translate", "outputs": [], "chapters": [{}]}
        report = monitor_performance.analyze(events, state, base + timedelta(seconds=15))
        first, active = report["attempts"]
        self.assertEqual((first["startup_seconds"], first["tool_seconds"],
                          first["runtime_non_tool_seconds"], first["post_turn_seconds"]), (1, 4, 4, 1))
        self.assertEqual(first["usage_updates"], 2)
        self.assertEqual(first["tokens"]["inputTokens"], 300)
        self.assertEqual(first["model"], "test-model")
        self.assertEqual(active["tokens"]["inputTokens"], 150)
        self.assertFalse(active["finished"])
        self.assertEqual(report["completed_attempt_totals"]["tokens"]["outputTokens"], 30)
        self.assertEqual(report["completed_attempt_totals"]["wall_seconds"], 10)


class CodexTransportTests(unittest.TestCase):
    SERVER = r'''
import json, sys
def send(value):
    print(json.dumps(value), flush=True)
for line in sys.stdin:
    m = json.loads(line)
    method = m.get("method")
    if method == "initialize":
        send({"id": m["id"], "result": {"userAgent": "fixture"}})
    elif method in ("thread/start", "thread/resume"):
        assert m["params"]["sandbox"] == "workspace-write"
        assert m["params"]["approvalPolicy"] == "never"
        send({"id": m["id"], "result": {"thread": {"id": "thread-1"}}})
    elif method == "turn/start":
        assert m["params"]["outputSchema"]["type"] == "object"
        assert not m["params"]["sandboxPolicy"]["networkAccess"]
        send({"id": 88, "method": "item/commandExecution/requestApproval",
              "params": {"command": "fixture denied command"}})
        answer = json.loads(sys.stdin.readline())
        assert answer["result"]["decision"] == "decline"
        # Notifications intentionally precede the RPC response.
        send({"method": "item/completed", "params": {"threadId": "thread-1", "turnId": "turn-1",
              "item": {"type": "agentMessage", "text": json.dumps({"outcome": "blocked"})}}})
        send({"method": "turn/completed", "params": {"threadId": "thread-1",
              "turn": {"id": "turn-1", "status": "completed"}}})
        send({"id": m["id"], "result": {"turn": {"id": "turn-1"}}})
'''

    def test_stdio_protocol_start_resume_interleaving_and_denial(self):
        with tempfile.TemporaryDirectory() as temp:
            for thread_id in (None, "thread-1"):
                with self.subTest(thread_id=thread_id):
                    events, threads = [], []
                    runtime = CodexRuntime([sys.executable, "-u", "-c", self.SERVER], timeout=10)
                    response = runtime.run(Request(Path(temp), "Fixture", None, thread_id, RESULT_SCHEMA),
                                           lambda *event: events.append(event), threads.append)
                    self.assertEqual(response, {"outcome": "blocked"})
                    self.assertEqual(threads, ["thread-1"])
                    self.assertTrue(any(event[0] == "runtime.request" for event in events))

    def test_timeout_terminates_process(self):
        with tempfile.TemporaryDirectory() as temp:
            runtime = CodexRuntime([sys.executable, "-u", "-c", "import time; time.sleep(30)"], timeout=0.1)
            with self.assertRaises(TimeoutError):
                runtime.run(Request(Path(temp), "", None, None, RESULT_SCHEMA), lambda *args: None, lambda t: None)


if __name__ == "__main__":
    unittest.main()

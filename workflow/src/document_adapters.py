"""Deterministic DOCX construction with interchangeable Word/LibreOffice renderers."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
import importlib.metadata
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Protocol

from .storage import digest

REPO = Path(__file__).resolve().parents[2]
DEFAULT_SETTINGS = {
    "schema_version": 1, "profile": "reader-facing", "body_font": "Georgia",
    "fallback_font": "Times New Roman", "font_size": 11, "title_font_size": 24,
    "heading_font_size": 16, "folio_font_size": 9, "toc_font_size": 11,
    "page_width_inches": 6, "page_height_inches": 9, "margin_inches": .75,
    "footer_inches": .35, "first_line_indent_inches": .2, "line_spacing": 1.2,
    "alignment": "left", "include_toc": True, "chapter_new_page": True,
    "indent_opening_paragraphs": False,
    "title_space_after_pt": 24, "heading_space_before_pt": 24,
    "heading_space_after_pt": 18, "toc_space_after_pt": 6,
}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def validate_settings(settings):
    require(isinstance(settings, dict) and set(settings) == set(DEFAULT_SETTINGS), "Incomplete/unknown format settings.")
    require(type(settings["schema_version"]) is int and settings["schema_version"] == 1, "Unsupported settings version.")
    require(settings["profile"] in ("submission", "editorial-review", "reader-facing"), "Invalid edition profile.")
    require(settings["alignment"] in ("left", "justify"), "Unsupported body alignment.")
    for key in ("body_font", "fallback_font"):
        require(isinstance(settings[key], str) and 0 < len(settings[key].strip()) <= 100, f"Invalid {key}.")
    for key in ("include_toc", "chapter_new_page", "indent_opening_paragraphs"):
        require(type(settings[key]) is bool, f"{key} must be boolean.")
    ranges = {
        "font_size": (8, 24), "title_font_size": (8, 60), "heading_font_size": (8, 36),
        "folio_font_size": (6, 18), "toc_font_size": (8, 24), "page_width_inches": (4, 14),
        "page_height_inches": (5, 20), "margin_inches": (.25, 3), "footer_inches": (.15, 2),
        "first_line_indent_inches": (0, 1), "line_spacing": (1, 3),
        "title_space_after_pt": (0, 72), "heading_space_before_pt": (0, 72),
        "heading_space_after_pt": (0, 72), "toc_space_after_pt": (0, 36),
    }
    for key, (lo, hi) in ranges.items():
        value = settings[key]
        require(type(value) in (int, float) and math.isfinite(value) and lo <= value <= hi, f"Invalid {key}.")
    require(settings["page_width_inches"] - 2 * settings["margin_inches"] >= 2,
            "Page margins leave insufficient text width.")
    require(settings["page_height_inches"] - 2 * settings["margin_inches"] >= 3
            and settings["footer_inches"] < settings["margin_inches"], "Invalid vertical page geometry.")
    return settings


@dataclass(frozen=True)
class RenderResult:
    pdf: Path
    renderer: str
    version: str


class Renderer(Protocol):
    def render(self, docx: Path, directory: Path) -> RenderResult: ...


def check_font(font, renderer):
    if renderer == "word":
        import winreg
        names = []
        for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            try:
                with winreg.OpenKey(hive, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts") as key:
                    for index in range(winreg.QueryInfoKey(key)[1]):
                        names.append(winreg.EnumValue(key, index)[0].lower())
            except FileNotFoundError:
                pass
        require(any(name == font.lower() or name.startswith(font.lower() + " ") for name in names),
                f"Requested font is not registered with Windows: {font}")
    else:
        executable = shutil.which("fc-match")
        require(executable, "Font verification requires fontconfig/fc-match in the LibreOffice environment.")
        result = subprocess.run([executable, "-f", "%{family}", font], capture_output=True,
                                text=True, encoding="utf-8", timeout=15, check=True)
        require(font.casefold() in [part.strip().casefold() for part in result.stdout.split(",")],
                f"Requested font would be substituted in the cloud renderer: {font}")


def word_worker(docx=None, pdf=None):
    """Run in a child process. DispatchEx obtains a dedicated Word instance."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    application, document, owned, old_update_links = None, None, False, None
    try:
        application = win32com.client.DispatchEx("Word.Application")
        require(application.Documents.Count == 0, "Word returned an instance with open user documents; refusing to use it.")
        owned = True
        application.Visible = False
        application.DisplayAlerts = 0
        application.AutomationSecurity = 3  # Disable macros for documents opened by this instance.
        old_update_links = application.Options.UpdateLinksAtOpen
        application.Options.UpdateLinksAtOpen = False
        version = str(application.Version)
        if docx is not None:
            document = application.Documents.Open(
                str(docx.resolve()), ConfirmConversions=False, ReadOnly=False,
                AddToRecentFiles=False, Visible=False,
            )
            document.Fields.Update()
            for section in document.Sections:
                for footer in section.Footers:
                    footer.Range.Fields.Update()
            document.Repaginate()
            document.Save()
            document.ExportAsFixedFormat(
                OutputFileName=str(pdf.resolve()), ExportFormat=17, OpenAfterExport=False,
                OptimizeFor=0, CreateBookmarks=1, DocStructureTags=True,
            )
        return version
    finally:
        try:
            if document is not None:
                document.Close(SaveChanges=0)
        finally:
            try:
                if owned and application is not None:
                    if old_update_links is not None:
                        application.Options.UpdateLinksAtOpen = old_update_links
                    if application.Documents.Count == 0:
                        application.Quit(SaveChanges=0)
            finally:
                document = application = None
                pythoncom.CoUninitialize()


class WordRenderer:
    def render(self, docx, directory):
        require(os.name == "nt", "Word renderer requires Windows with Microsoft Word.")
        pdf = directory / (docx.stem + ".pdf")
        result = subprocess.run(
            [sys.executable, "-m", "workflow.src.document_adapters", "word-render", str(docx), str(pdf)],
            cwd=REPO, capture_output=True, text=True, encoding="utf-8", timeout=180,
        )
        require(result.returncode == 0, "Word rendering failed: " + result.stderr[-3000:])
        require(pdf.is_file() and pdf.stat().st_size > 0, "Word did not produce a PDF.")
        return RenderResult(pdf, "word", json.loads(result.stdout)["version"])


class LibreOfficeRenderer:
    def render(self, docx, directory):
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        require(executable, "LibreOffice renderer is not installed or not on PATH.")
        pdf = directory / (docx.stem + ".pdf")
        with tempfile.TemporaryDirectory(prefix="manuscript-lo-") as profile:
            result = subprocess.run([
                executable, "-env:UserInstallation=" + Path(profile).as_uri(),
                "--headless", "--convert-to", "pdf:writer_pdf_Export",
                "--outdir", str(directory), str(docx),
            ], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180)
            require(result.returncode == 0 and pdf.is_file() and pdf.stat().st_size > 0,
                    "LibreOffice rendering failed: " + result.stderr[-3000:] + result.stdout[-1000:])
        version = subprocess.run([executable, "--version"], capture_output=True, text=True,
                                 encoding="utf-8", timeout=15, check=True).stdout.strip()
        return RenderResult(pdf, "libreoffice", version)


def preflight(renderer):
    require(renderer in ("word", "libreoffice"), "Unknown document renderer.")
    for package in ("python-docx", "markdown-it-py", "pypdfium2", "Pillow"):
        importlib.metadata.version(package)
    if renderer == "word":
        require(os.name == "nt", "Word renderer requires Windows.")
        result = subprocess.run([sys.executable, "-m", "workflow.src.document_adapters", "word-probe"],
                                cwd=REPO, capture_output=True, text=True, encoding="utf-8", timeout=45)
        require(result.returncode == 0, "Word automation unavailable: " + result.stderr[-2000:])
        return {"renderer": renderer, "version": json.loads(result.stdout)["version"]}
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    require(executable, "LibreOffice renderer is unavailable.")
    return {"renderer": renderer, "executable": executable}


def content_blocks(markdown):
    from markdown_it import MarkdownIt
    blocks, heading, quote = [], None, 0
    for token in MarkdownIt("commonmark").parse(markdown):
        if token.type in ("html_block", "html_inline", "image", "hr", "fence", "code_block",
                          "bullet_list_open", "ordered_list_open"):
            raise ValueError(f"Unsupported manuscript block for deterministic DOCX construction: {token.type}")
        if token.type == "blockquote_open":
            quote += 1
        elif token.type == "blockquote_close":
            quote -= 1
        elif token.type == "heading_open":
            heading = int(token.tag[1:])
        elif token.type == "heading_close":
            heading = None
        elif token.type == "inline":
            require(not any(child.type in ("image", "html_inline") for child in token.children or []),
                    "Unsupported manuscript block: inline image/HTML")
            blocks.append({"heading": heading, "quote": quote > 0, "tokens": token.children or []})
    return blocks


class PythonDocxBuilder:
    def build(self, markdown, settings, path, toc_pages):
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
        from docx.opc.constants import RELATIONSHIP_TYPE

        validate_settings(settings)
        document = Document()
        document.core_properties.author = ""
        document.core_properties.last_modified_by = ""
        document.core_properties.comments = ""
        section = document.sections[0]
        section.page_width, section.page_height = Inches(settings["page_width_inches"]), Inches(settings["page_height_inches"])
        section.top_margin = section.bottom_margin = Inches(settings["margin_inches"])
        section.left_margin = section.right_margin = Inches(settings["margin_inches"])
        section.footer_distance = Inches(settings["footer_inches"])
        controlled_styles = []
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "TOC Heading", "TOC 1", "TOC 2", "TOC 3"):
            style = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            controlled_styles.append(style)
            style.font.name = settings["body_font"]
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.size = Pt(settings["font_size"])
            style.font.bold = style.font.italic = False
            if name != "Normal":
                style.paragraph_format.line_spacing = 1
                if not name.startswith("Heading "):
                    style.paragraph_format.page_break_before = False
            ppr = style.element.get_or_add_pPr()
            outline = ppr.find(qn("w:outlineLvl"))
            if outline is None:
                outline = OxmlElement("w:outlineLvl")
                ppr.append(outline)
            outline.set(qn("w:val"), str(int(name[-1]) - 1) if name.startswith("Heading ") else "9")
        # Word theme fonts outrank explicit ascii/hAnsi names. Remove them from
        # both paragraph styles and their linked character styles before saving.
        linked_ids = {node.get(qn("w:val")) for style in controlled_styles
                      for node in style.element.findall(qn("w:link"))}
        controlled_styles.extend(style for style in document.styles if style.style_id in linked_ids)
        for style in controlled_styles:
            rpr = style.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.insert(0, fonts)
            for attribute in list(fonts.attrib):
                if "theme" in attribute.lower():
                    del fonts.attrib[attribute]
            for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + slot), settings["body_font"])
            style.font.color.rgb = RGBColor(0, 0, 0)
            color = rpr.find(qn("w:color"))
            for attribute in list(color.attrib):
                if "theme" in attribute.lower():
                    del color.attrib[attribute]
            for node in list(rpr):
                if node.tag in (qn("w:spacing"), qn("w:bCs"), qn("w:iCs")):
                    rpr.remove(node)
            ppr = style.element.find(qn("w:pPr"))
            if ppr is not None:
                for border in list(ppr.findall(qn("w:pBdr"))):
                    ppr.remove(border)
        normal = document.styles["Normal"].paragraph_format
        normal.line_spacing = settings["line_spacing"]
        normal.space_before = normal.space_after = Pt(0)
        normal.first_line_indent = Inches(settings["first_line_indent_inches"])
        normal.widow_control = True
        normal.alignment = WD_ALIGN_PARAGRAPH.LEFT if settings["alignment"] == "left" else WD_ALIGN_PARAGRAPH.JUSTIFY
        title = document.styles["Title"]
        title.font.size = Pt(settings["title_font_size"])
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.first_line_indent = Inches(0)
        title.paragraph_format.space_after = Pt(settings["title_space_after_pt"])
        toc_heading = document.styles["TOC Heading"]
        toc_heading.font.size = Pt(max(settings["font_size"], settings["heading_font_size"] - 2))
        toc_heading.paragraph_format.space_before = Pt(12)
        toc_heading.paragraph_format.space_after = Pt(12)
        for level in range(1, 4):
            style = document.styles[f"Heading {level}"]
            style.font.size = Pt(max(settings["font_size"], settings["heading_font_size"] - 2 * (level - 1)))
            style.paragraph_format.first_line_indent = Inches(0)
            style.paragraph_format.space_before = Pt(settings["heading_space_before_pt"])
            style.paragraph_format.space_after = Pt(settings["heading_space_after_pt"])
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style.paragraph_format.page_break_before = settings["chapter_new_page"] and level == 1
        blocks = content_blocks(markdown)
        require(blocks, "Manuscript is empty.")
        headings = []
        for index, block in enumerate(blocks):
            if block["heading"] and index != 0:
                text = "".join(t.content for t in block["tokens"] if t.type in ("text", "code_inline", "text_special"))
                headings.append({"source_index": index, "text": text, "bookmark": f"section_{index}",
                                 "level": min(3, max(1, block["heading"] - 1))})
        by_index = {h["source_index"]: h for h in headings}
        document.core_properties.title = "".join(t.content for t in blocks[0]["tokens"]) if blocks[0]["heading"] else ""

        def inline(paragraph, tokens):
            bold, italic, link = False, False, None
            for token in tokens:
                if token.type in ("strong_open", "strong_close"):
                    bold = token.type.endswith("_open")
                elif token.type in ("em_open", "em_close"):
                    italic = token.type.endswith("_open")
                elif token.type == "link_open":
                    url = token.attrGet("href")
                    require(url and re.match(r"^(https?://|mailto:|#)", url), "Unsupported manuscript hyperlink target.")
                    link = OxmlElement("w:hyperlink")
                    if url.startswith("#"):
                        link.set(qn("w:anchor"), url[1:])
                    else:
                        link.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                    paragraph._p.append(link)
                elif token.type == "link_close":
                    link = None
                elif token.type in ("text", "code_inline", "text_special", "softbreak", "hardbreak"):
                    run = paragraph.add_run()
                    run.bold, run.italic = bold, italic
                    if token.type == "hardbreak":
                        run.add_break(WD_BREAK.LINE)
                    else:
                        run.text = " " if token.type == "softbreak" else token.content
                    if token.type == "code_inline":
                        run.font.name = "Consolas"
                    if link is not None:
                        link.append(run._r)
                else:
                    raise ValueError(f"Unsupported inline element: {token.type}")

        for index, block in enumerate(blocks):
            is_title = index == 0 and block["heading"] is not None
            style = "Title" if is_title else f"Heading {by_index[index]['level']}" if index in by_index else "Normal"
            paragraph = document.add_paragraph(style=style)
            if index in by_index:
                start = OxmlElement("w:bookmarkStart")
                start.set(qn("w:id"), str(index))
                start.set(qn("w:name"), by_index[index]["bookmark"])
                paragraph._p.append(start)
            inline(paragraph, block["tokens"])
            if index in by_index:
                end = OxmlElement("w:bookmarkEnd")
                end.set(qn("w:id"), str(index))
                paragraph._p.append(end)
            if not settings["indent_opening_paragraphs"] and not block["heading"] and (index == 0 or blocks[index - 1]["heading"]):
                paragraph.paragraph_format.first_line_indent = Inches(0)
            if block["quote"]:
                paragraph.paragraph_format.left_indent = Inches(.2)
                paragraph.paragraph_format.right_indent = Inches(.2)
            if is_title and settings["include_toc"] and headings:
                label = document.add_paragraph("Contents", style="TOC Heading")
                label.paragraph_format.first_line_indent = Inches(0)
                for item in headings:
                    toc = document.add_paragraph(style=f"TOC {item['level']}")
                    toc.paragraph_format.first_line_indent = Inches(0)
                    toc.paragraph_format.space_after = Pt(settings["toc_space_after_pt"])
                    toc.paragraph_format.tab_stops.add_tab_stop(
                        Inches(settings["page_width_inches"] - 2 * settings["margin_inches"]),
                        WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS,
                    )
                    hyperlink = OxmlElement("w:hyperlink")
                    hyperlink.set(qn("w:anchor"), item["bookmark"])
                    run = toc.add_run(item["text"] + "\t" + str(toc_pages.get(item["source_index"], 0)))
                    run.font.size = Pt(settings["toc_font_size"])
                    hyperlink.append(run._r)
                    toc._p.append(hyperlink)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent = Inches(0)
        run = footer.add_run()
        run.font.name, run.font.size = settings["body_font"], Pt(settings["folio_font_size"])
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run.text = "1"
        field.append(run._r)
        footer._p.append(field)
        hyphenation = OxmlElement("w:autoHyphenation")
        hyphenation.set(qn("w:val"), "0")
        document.settings.element.append(hyphenation)
        document.save(path)
        return headings


def pdf_outline_pages(pdf_path, headings):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(pdf_path)
    try:
        outline = list(pdf.get_toc())
        result, cursor = {}, 0
        for heading in headings:
            while cursor < len(outline) and outline[cursor].get_title().strip() != heading["text"].strip():
                cursor += 1
            destination = outline[cursor].get_dest() if cursor < len(outline) else None
            require(destination is not None,
                    f"Rendered PDF has no heading destination: {heading['text']}")
            result[heading["source_index"]] = destination.get_index() + 1
            cursor += 1
        return result
    finally:
        pdf.close()


def build_and_render(markdown, settings, directory, renderer_name, *, renderer=None):
    import pypdfium2 as pdfium
    validate_settings(settings)
    check_font(settings["body_font"], renderer_name)
    directory.mkdir(parents=True, exist_ok=False)
    renderer = renderer or (WordRenderer() if renderer_name == "word" else LibreOfficeRenderer())
    builder = PythonDocxBuilder()
    pages = {}
    for iteration in range(1, 5):
        attempt = directory / f"layout-{iteration}"
        attempt.mkdir()
        docx = attempt / "manuscript.docx"
        headings = builder.build(markdown, settings, docx, pages)
        rendered = renderer.render(docx, attempt)
        actual_pages = pdf_outline_pages(rendered.pdf, headings) if settings["include_toc"] else {}
        if actual_pages == pages:
            break
        pages = actual_pages
    else:
        raise ValueError("TOC pagination did not converge; revise the format specification.")
    pdf = pdfium.PdfDocument(rendered.pdf)
    images = []
    try:
        # Ensure the static TOC's displayed numbers actually survived rendering.
        if settings["include_toc"] and headings:
            toc_end = min(actual_pages.values()) - 1
            toc_text = ""
            for index in range(toc_end):
                page = pdf[index]
                textpage = page.get_textpage()
                toc_text += textpage.get_text_range() + "\n"
                textpage.close()
                page.close()
            for heading in headings:
                require(re.search(re.escape(heading["text"]) + r"[\s.·…]+" +
                                  str(actual_pages[heading["source_index"]]) + r"(?!\d)", toc_text),
                        f"Rendered TOC entry/page number is missing: {heading['text']}")
        image_dir = directory / "pages"
        image_dir.mkdir()
        for index in range(len(pdf)):
            page = pdf[index]
            bitmap = page.render(scale=1.5)
            image = bitmap.to_pil()
            path = image_dir / f"page-{index + 1:03d}.png"
            image.save(path)
            images.append({"page": index + 1, "path": str(path), "sha256": digest(path.read_bytes())})
            image.close()
            bitmap.close()
            page.close()
    finally:
        pdf.close()
    return {"docx": str(docx), "pdf": str(rendered.pdf), "pages": images,
            "renderer": rendered.renderer, "renderer_version": rendered.version,
            "builder": "python-docx", "builder_version": importlib.metadata.version("python-docx"),
            "toc_pages": pages, "layout_passes": iteration}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("operation", choices=("word-probe", "word-render"))
    parser.add_argument("docx", type=Path, nargs="?")
    parser.add_argument("pdf", type=Path, nargs="?")
    args = parser.parse_args()
    if args.operation == "word-render":
        require(args.docx is not None and args.pdf is not None, "Word render needs DOCX and PDF paths.")
    print(json.dumps({"version": word_worker(args.docx, args.pdf)}))


if __name__ == "__main__":
    main()
